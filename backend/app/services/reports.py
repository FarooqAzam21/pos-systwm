from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import String, and_, case, cast, func, or_, select
from sqlalchemy.orm import Session

from app.models.audit_log import AuditLog
from app.models.business import Business
from app.models.category import Category
from app.models.customer import Customer
from app.models.daily_closing import DailyClosing
from app.models.expense import Expense, ExpenseCategory
from app.models.payment import Payment, PaymentReferenceType, PaymentType
from app.models.product import Product
from app.models.purchase import Purchase
from app.models.sale import Sale, SaleItem, SaleStatus
from app.models.sale_return import SaleReturn, SaleReturnItem
from app.models.stock_movement import StockMovement
from app.models.supplier import Supplier
from app.models.user import User
from app.services.commerce import allocate_amount, money
from app.services.dates import business_datetime_bounds, business_local_date_expression

ZERO = Decimal("0")
MAX_RANGE_DAYS = 366
REPORT_FIELDS = {
    "sales": ["invoice", "date", "customer", "cashier", "subtotal", "discount", "tax", "total", "paid", "outstanding", "cogs", "gross_profit", "status"],
    "purchases": ["invoice", "date", "supplier", "subtotal", "discount", "tax", "total", "paid", "outstanding", "payment_status"],
    "expenses": ["date", "category", "description", "amount", "payment_method", "created_by"],
    "inventory": ["sku", "barcode", "product", "category", "current_stock", "cost_price", "selling_price", "inventory_value", "minimum_stock", "stock_status"],
    "stock-movements": ["date", "product", "movement_type", "quantity", "previous_stock", "new_stock", "unit_cost", "reference", "user"],
    "product-profitability": ["product", "units_sold", "revenue", "cogs", "gross_profit", "margin_percent"],
    "category-performance": ["category", "units_sold", "revenue", "cogs", "gross_profit", "margin_percent"],
    "customers": ["customer", "total_purchases", "total_paid", "outstanding", "transactions", "returns", "net_spending", "last_purchase"],
    "suppliers": ["supplier", "purchases", "paid", "outstanding", "purchase_count", "latest_purchase"],
    "customer-analytics": ["customer", "total_purchases", "total_paid", "outstanding", "transactions", "returns", "net_spending", "last_purchase"],
    "supplier-analytics": ["supplier", "purchases", "paid", "outstanding", "purchase_count", "latest_purchase"],
    "payments": ["date", "reference", "party", "type", "reference_type", "amount", "payment_method", "received_by"],
}


def normalize_range(start_date: date | None, end_date: date | None) -> tuple[date, date]:
    end = end_date or date.today()
    start = start_date or end - timedelta(days=29)
    if end < start:
        raise ValueError("End date must not be before start date")
    if (end - start).days > MAX_RANGE_DAYS:
        raise ValueError(f"Date range cannot exceed {MAX_RANGE_DAYS} days")
    return start, end


def _date_filter(model, start: date, end: date):
    return and_(model >= start, model <= end)


def _serialize_date(value):
    return value.isoformat() if value else None


def _percent(profit: Decimal, revenue: Decimal) -> Decimal:
    return money(profit / revenue * Decimal(100)) if revenue else ZERO


def _filter_sales(db, stmt, business_id, start, end, filters):
    start_at, end_at = business_datetime_bounds(db, business_id, start, end)
    stmt = stmt.where(Sale.business_id == business_id, Sale.sale_date >= start_at, Sale.sale_date < end_at)
    if filters.get("customer_id"):
        stmt = stmt.where(Sale.customer_id == filters["customer_id"])
    if filters.get("cashier_id"):
        stmt = stmt.where(Sale.created_by == filters["cashier_id"])
    if filters.get("status"):
        stmt = stmt.where(Sale.status == filters["status"])
    if filters.get("payment_status"):
        stmt = stmt.where(Sale.payment_status == filters["payment_status"])
    if filters.get("payment_method"):
        payment_exists = select(Payment.id).where(
            Payment.business_id == business_id,
            Payment.reference_type == PaymentReferenceType.SALE,
            Payment.reference_id == cast(Sale.id, String),
            Payment.payment_method == filters["payment_method"],
        ).exists()
        stmt = stmt.where(payment_exists)
    if filters.get("search"):
        term = f"%{filters['search'].strip()}%"
        matching_customers = select(Customer.id).where(
            Customer.business_id == business_id,
            Customer.name.ilike(term),
        )
        stmt = stmt.where(
            or_(
                Sale.invoice_number.ilike(term),
                Sale.customer_id.in_(matching_customers),
            )
        )
    return stmt


def _profitability_rows(db: Session, business_id, start, end, by_category=False, category_id=None):
    start_at, end_at = business_datetime_bounds(db, business_id, start, end)
    joined = db.execute(
        select(SaleItem, Sale, Product.id, Product.name, Product.sku, Product.category_id, Category.name)
        .join(Sale, and_(Sale.id == SaleItem.sale_id, Sale.business_id == business_id))
        .join(Product, and_(Product.id == SaleItem.product_id, Product.business_id == business_id))
        .outerjoin(Category, and_(Category.id == Product.category_id, Category.business_id == business_id))
        .where(
            SaleItem.business_id == business_id,
            Sale.sale_date >= start_at,
            Sale.sale_date < end_at,
            Sale.status != SaleStatus.VOIDED,
        )
    ).all()
    by_sale = defaultdict(list)
    for row in joined:
        by_sale[row[1].id].append(row)
    grouped = {}

    def add_values(product_id, product_name, sku, product_category_id, category_name, quantity, revenue, cogs):
        if category_id and product_category_id != category_id:
            return
        group_key = (product_category_id, category_name or "Uncategorized") if by_category else product_id
        target = grouped.setdefault(group_key, {
            "product": product_name if not by_category else category_name or "Uncategorized",
            "sku": sku if not by_category else None,
            "category_id": product_category_id,
            "category": category_name or "Uncategorized",
            "units_sold": ZERO,
            "revenue": ZERO,
            "cogs": ZERO,
        })
        target["units_sold"] += quantity
        target["revenue"] += revenue
        target["cogs"] += cogs

    for rows in by_sale.values():
        sale = rows[0][1]
        global_discount = Decimal(sale.discount) - sum((Decimal(row[0].discount) for row in rows), ZERO)
        global_tax = Decimal(sale.tax) - sum((Decimal(row[0].tax) for row in rows), ZERO)
        if global_discount < ZERO or global_tax < ZERO:
            continue
        discount_weights = [
            max(ZERO, money(Decimal(row[0].quantity) * Decimal(row[0].unit_price)) - Decimal(row[0].discount))
            for row in rows
        ]
        tax_weights = [money(Decimal(row[0].quantity) * Decimal(row[0].unit_price)) for row in rows]
        if sum(tax_weights, ZERO) == ZERO:
            tax_weights = [Decimal(row[0].quantity) for row in rows]
        allocated_discounts = allocate_amount(global_discount, discount_weights)
        allocated_taxes = allocate_amount(global_tax, tax_weights)
        for index, row in enumerate(rows):
            item, _, product_id, product_name, sku, product_category_id, category_name = row
            revenue = money(Decimal(item.line_total) - allocated_discounts[index] + allocated_taxes[index])
            add_values(
                product_id, product_name, sku, product_category_id, category_name,
                Decimal(item.quantity), revenue, Decimal(item.cogs),
            )

    returns = db.execute(
        select(
            Product.id,
            Product.name,
            Product.sku,
            Product.category_id,
            Category.name,
            SaleReturnItem.quantity,
            SaleReturnItem.refund_amount,
            SaleReturnItem.cogs,
        )
        .join(SaleReturn, and_(SaleReturn.id == SaleReturnItem.return_id, SaleReturn.business_id == business_id))
        .join(Product, and_(Product.id == SaleReturnItem.product_id, Product.business_id == business_id))
        .outerjoin(Category, and_(Category.id == Product.category_id, Category.business_id == business_id))
        .where(
            SaleReturnItem.business_id == business_id,
            SaleReturn.created_at >= start_at,
            SaleReturn.created_at < end_at,
        )
    ).all()
    for product_id, product_name, sku, product_category_id, category_name, quantity, refund, cogs in returns:
        add_values(product_id, product_name, sku, product_category_id, category_name, -Decimal(quantity), -Decimal(refund), -Decimal(cogs))
    output = []
    for target in grouped.values():
        target["revenue"] = money(target["revenue"])
        target["cogs"] = money(target["cogs"])
        target["gross_profit"] = money(target["revenue"] - target["cogs"])
        target["margin_percent"] = _percent(target["gross_profit"], target["revenue"])
        output.append(target)
    return sorted(output, key=lambda item: item["revenue"], reverse=True)


def get_report_data(
    db: Session,
    business_id,
    report_name: str,
    start: date,
    end: date,
    filters: dict | None = None,
    page: int = 1,
    limit: int = 50,
) -> dict:
    filters = filters or {}
    if report_name not in REPORT_FIELDS:
        raise ValueError("Unknown report")

    if report_name == "sales":
        base = _filter_sales(db, select(Sale.id), business_id, start, end, filters).outerjoin(
            Customer, and_(Customer.id == Sale.customer_id, Customer.business_id == business_id)
        )
        statement = (
            _filter_sales(
                db,
                select(Sale.invoice_number.label("invoice"), Sale.sale_date.label("date"),
                       Customer.name.label("customer"), User.name.label("cashier"),
                       Sale.subtotal.label("subtotal"), Sale.discount.label("discount"), Sale.tax.label("tax"),
                       Sale.total_amount.label("total"), Sale.paid_amount.label("paid"),
                       Sale.outstanding_amount.label("outstanding"), Sale.cogs_amount.label("cogs"),
                       Sale.gross_profit.label("gross_profit"), Sale.status.label("status")),
                business_id, start, end, filters,
            )
            .outerjoin(Customer, and_(Customer.id == Sale.customer_id, Customer.business_id == business_id))
            .outerjoin(User, and_(User.id == Sale.created_by, User.business_id == business_id))
        )
        total = db.scalar(select(func.count()).select_from(base.subquery())) or 0
        numeric = {"subtotal": func.sum(Sale.subtotal), "discount": func.sum(Sale.discount), "tax": func.sum(Sale.tax), "total": func.sum(Sale.total_amount), "paid": func.sum(Sale.paid_amount), "outstanding": func.sum(Sale.outstanding_amount), "cogs": func.sum(Sale.cogs_amount), "gross_profit": func.sum(Sale.gross_profit)}
        summary = db.execute(_filter_sales(db, select(*numeric.values()), business_id, start, end, filters).outerjoin(Customer, and_(Customer.id == Sale.customer_id, Customer.business_id == business_id))).one()
        rows = db.execute(statement.order_by(Sale.sale_date.desc()).offset((page - 1) * limit).limit(limit)).mappings().all()
        return {"items": [dict(row) for row in rows], "total": total, "summary": dict(zip(numeric.keys(), summary))}

    if report_name == "purchases":
        statement = select(
            Purchase.invoice_number.label("invoice"), Purchase.purchase_date.label("date"),
            Supplier.name.label("supplier"), Purchase.subtotal.label("subtotal"),
            Purchase.discount.label("discount"), Purchase.tax.label("tax"),
            Purchase.total_amount.label("total"), Purchase.paid_amount.label("paid"),
            Purchase.outstanding_amount.label("outstanding"), Purchase.payment_status.label("payment_status"),
        ).join(Supplier, and_(Supplier.id == Purchase.supplier_id, Supplier.business_id == business_id)).where(
            Purchase.business_id == business_id, _date_filter(Purchase.purchase_date, start, end),
        )
        if filters.get("supplier_id"):
            statement = statement.where(Purchase.supplier_id == filters["supplier_id"])
        if filters.get("payment_status"):
            statement = statement.where(Purchase.payment_status == filters["payment_status"])
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        summary = db.execute(
            select(
                func.coalesce(func.sum(Purchase.total_amount), 0),
                func.coalesce(func.sum(Purchase.paid_amount), 0),
                func.coalesce(func.sum(Purchase.outstanding_amount), 0),
            )
            .select_from(Purchase)
            .join(Supplier, and_(Supplier.id == Purchase.supplier_id, Supplier.business_id == business_id))
            .where(Purchase.business_id == business_id, _date_filter(Purchase.purchase_date, start, end))
            .where(*([Purchase.supplier_id == filters["supplier_id"]] if filters.get("supplier_id") else []))
            .where(*([Purchase.payment_status == filters["payment_status"]] if filters.get("payment_status") else []))
        ).one()
        rows = db.execute(statement.order_by(Purchase.purchase_date.desc()).offset((page - 1) * limit).limit(limit)).mappings().all()
        return {"items": [dict(row) for row in rows], "total": total, "summary": {"total": summary[0] or ZERO, "paid": summary[1] or ZERO, "outstanding": summary[2] or ZERO}}

    if report_name == "expenses":
        statement = select(
            Expense.expense_date.label("date"), ExpenseCategory.name.label("category"),
            Expense.description.label("description"), Expense.amount.label("amount"),
            Expense.payment_method.label("payment_method"), User.name.label("created_by"),
        ).join(ExpenseCategory, and_(ExpenseCategory.id == Expense.category_id, ExpenseCategory.business_id == business_id)).outerjoin(
            User, and_(User.id == Expense.created_by, User.business_id == business_id)
        ).where(Expense.business_id == business_id, _date_filter(Expense.expense_date, start, end))
        if filters.get("category_id"):
            statement = statement.where(Expense.category_id == filters["category_id"])
        if filters.get("payment_method"):
            statement = statement.where(Expense.payment_method == filters["payment_method"])
        if filters.get("search"):
            statement = statement.where(Expense.description.ilike(f"%{filters['search'].strip()}%"))
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        summary_stmt = select(func.coalesce(func.sum(Expense.amount), 0)).where(
            Expense.business_id == business_id,
            _date_filter(Expense.expense_date, start, end),
        )
        if filters.get("category_id"):
            summary_stmt = summary_stmt.where(Expense.category_id == filters["category_id"])
        if filters.get("payment_method"):
            summary_stmt = summary_stmt.where(Expense.payment_method == filters["payment_method"])
        if filters.get("search"):
            summary_stmt = summary_stmt.where(Expense.description.ilike(f"%{filters['search'].strip()}%"))
        summary = db.scalar(summary_stmt) or ZERO
        rows = db.execute(statement.order_by(Expense.expense_date.desc(), Expense.created_at.desc()).offset((page - 1) * limit).limit(limit)).mappings().all()
        return {"items": [dict(row) for row in rows], "total": total, "summary": {"amount": summary}}

    if report_name == "inventory":
        value = Product.current_stock * Product.cost_price
        statement = select(
            Product.sku.label("sku"), Product.barcode.label("barcode"), Product.name.label("product"),
            Category.name.label("category"), Product.current_stock.label("current_stock"),
            Product.cost_price.label("cost_price"), Product.selling_price.label("selling_price"),
            value.label("inventory_value"), Product.minimum_stock.label("minimum_stock"),
            case(
                (Product.current_stock <= 0, "OUT_OF_STOCK"),
                (Product.current_stock <= Product.minimum_stock, "LOW_STOCK"),
                else_="IN_STOCK",
            ).label("stock_status"),
        ).outerjoin(Category, and_(Category.id == Product.category_id, Category.business_id == business_id)).where(
            Product.business_id == business_id, Product.is_active.is_(True),
        )
        if filters.get("category_id"):
            statement = statement.where(Product.category_id == filters["category_id"])
        if filters.get("search"):
            term = f"%{filters['search'].strip()}%"
            statement = statement.where(or_(Product.name.ilike(term), Product.sku.ilike(term), Product.barcode.ilike(term)))
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        valuation_stmt = select(
            func.coalesce(func.sum(Product.current_stock * Product.cost_price), 0),
            func.coalesce(func.sum(Product.current_stock * Product.selling_price), 0),
        ).where(Product.business_id == business_id, Product.is_active.is_(True))
        if filters.get("category_id"):
            valuation_stmt = valuation_stmt.where(Product.category_id == filters["category_id"])
        if filters.get("search"):
            term = f"%{filters['search'].strip()}%"
            valuation_stmt = valuation_stmt.where(
                or_(Product.name.ilike(term), Product.sku.ilike(term), Product.barcode.ilike(term))
            )
        valuation = db.execute(valuation_stmt).one()
        rows = db.execute(statement.order_by(Product.name).offset((page - 1) * limit).limit(limit)).mappings().all()
        cost_value, retail_value = valuation
        return {"items": [dict(row) for row in rows], "total": total, "summary": {
            "inventory_cost_value": cost_value, "potential_retail_value": retail_value,
            "potential_gross_profit": (retail_value or ZERO) - (cost_value or ZERO),
        }}

    if report_name == "stock-movements":
        start_at, end_at = business_datetime_bounds(db, business_id, start, end)
        statement = select(
            StockMovement.created_at.label("date"), Product.name.label("product"),
            StockMovement.movement_type.label("movement_type"), StockMovement.quantity.label("quantity"),
            StockMovement.previous_stock.label("previous_stock"), StockMovement.new_stock.label("new_stock"),
            StockMovement.unit_cost.label("unit_cost"),
            func.coalesce(StockMovement.reference_type, "") + ":" + func.coalesce(StockMovement.reference_id, "").label("reference"),
            User.name.label("user"),
        ).join(Product, and_(Product.id == StockMovement.product_id, Product.business_id == business_id)).outerjoin(
            User, and_(User.id == StockMovement.created_by, User.business_id == business_id)
        ).where(StockMovement.business_id == business_id, StockMovement.created_at >= start_at, StockMovement.created_at < end_at)
        if filters.get("product_id"):
            statement = statement.where(StockMovement.product_id == filters["product_id"])
        if filters.get("movement_type"):
            statement = statement.where(StockMovement.movement_type == filters["movement_type"])
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        rows = db.execute(statement.order_by(StockMovement.created_at.desc()).offset((page - 1) * limit).limit(limit)).mappings().all()
        return {"items": [dict(row) for row in rows], "total": total, "summary": {}}

    if report_name in {"product-profitability", "category-performance"}:
        rows = _profitability_rows(
            db, business_id, start, end,
            by_category=report_name == "category-performance",
            category_id=filters.get("category_id"),
        )
        total = len(rows)
        selected = rows[(page - 1) * limit: page * limit]
        return {"items": selected, "total": total, "summary": {
            "units_sold": sum((row["units_sold"] for row in rows), ZERO),
            "revenue": money(sum((row["revenue"] for row in rows), ZERO)),
            "cogs": money(sum((row["cogs"] for row in rows), ZERO)),
            "gross_profit": money(sum((row["gross_profit"] for row in rows), ZERO)),
        }}

    if report_name in {"customers", "customer-analytics"}:
        start_at, end_at = business_datetime_bounds(db, business_id, start, end)
        sales_sub = select(
            Sale.customer_id.label("customer_id"),
            func.sum(Sale.subtotal - Sale.discount + Sale.tax).label("total_purchases"),
            func.count(Sale.id).label("transactions"),
            func.max(Sale.sale_date).label("last_purchase"),
        ).where(Sale.business_id == business_id, Sale.customer_id.is_not(None), Sale.sale_date >= start_at, Sale.sale_date < end_at, Sale.status != SaleStatus.VOIDED).group_by(Sale.customer_id).subquery()
        paid_sub = select(Payment.customer_id.label("customer_id"), func.sum(Payment.amount).label("total_paid")).where(
            Payment.business_id == business_id, Payment.payment_type == PaymentType.INCOMING,
            Payment.customer_id.is_not(None), Payment.payment_date >= start_at, Payment.payment_date < end_at,
        ).group_by(Payment.customer_id).subquery()
        return_sub = select(
            Sale.customer_id.label("customer_id"),
            func.sum(SaleReturn.refund_amount).label("returns"),
        ).join(Sale, and_(Sale.id == SaleReturn.sale_id, Sale.business_id == business_id)).where(
            SaleReturn.business_id == business_id,
            SaleReturn.created_at >= start_at,
            SaleReturn.created_at < end_at,
        ).group_by(Sale.customer_id).subquery()
        total_purchases = func.coalesce(sales_sub.c.total_purchases, 0)
        total_returns = func.coalesce(return_sub.c.returns, 0)
        statement = select(
            Customer.name.label("customer"),
            total_purchases.label("total_purchases"),
            func.coalesce(paid_sub.c.total_paid, 0).label("total_paid"),
            Customer.current_balance.label("outstanding"),
            func.coalesce(sales_sub.c.transactions, 0).label("transactions"),
            total_returns.label("returns"),
            (total_purchases - total_returns).label("net_spending"),
            sales_sub.c.last_purchase.label("last_purchase"),
        ).where(Customer.business_id == business_id).outerjoin(sales_sub, sales_sub.c.customer_id == Customer.id).outerjoin(
            paid_sub, paid_sub.c.customer_id == Customer.id
        ).outerjoin(return_sub, return_sub.c.customer_id == Customer.id)
        if filters.get("search"):
            statement = statement.where(Customer.name.ilike(f"%{filters['search'].strip()}%"))
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        rows = db.execute(statement.order_by(func.coalesce(sales_sub.c.total_purchases, 0).desc()).offset((page - 1) * limit).limit(limit)).mappings().all()
        return {"items": [dict(row) for row in rows], "total": total, "summary": {}}

    if report_name in {"suppliers", "supplier-analytics"}:
        purchase_sub = select(
            Purchase.supplier_id.label("supplier_id"),
            func.sum(Purchase.total_amount).label("purchases"),
            func.sum(Purchase.paid_amount).label("paid"),
            func.sum(Purchase.outstanding_amount).label("outstanding"),
            func.count(Purchase.id).label("purchase_count"),
            func.max(Purchase.purchase_date).label("latest_purchase"),
        ).where(Purchase.business_id == business_id, _date_filter(Purchase.purchase_date, start, end)).group_by(Purchase.supplier_id).subquery()
        statement = select(
            Supplier.name.label("supplier"), func.coalesce(purchase_sub.c.purchases, 0).label("purchases"),
            func.coalesce(purchase_sub.c.paid, 0).label("paid"),
            func.coalesce(purchase_sub.c.outstanding, 0).label("outstanding"),
            func.coalesce(purchase_sub.c.purchase_count, 0).label("purchase_count"),
            purchase_sub.c.latest_purchase.label("latest_purchase"),
        ).where(Supplier.business_id == business_id).outerjoin(purchase_sub, purchase_sub.c.supplier_id == Supplier.id)
        if filters.get("search"):
            statement = statement.where(Supplier.name.ilike(f"%{filters['search'].strip()}%"))
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        rows = db.execute(statement.order_by(func.coalesce(purchase_sub.c.purchases, 0).desc()).offset((page - 1) * limit).limit(limit)).mappings().all()
        return {"items": [dict(row) for row in rows], "total": total, "summary": {}}

    if report_name == "payments":
        start_at, end_at = business_datetime_bounds(db, business_id, start, end)
        statement = select(
            Payment.payment_date.label("date"),
            func.coalesce(Sale.invoice_number, Purchase.invoice_number, Payment.reference_id).label("reference"),
            func.coalesce(Customer.name, Supplier.name, "—").label("party"),
            Payment.payment_type.label("type"), Payment.reference_type.label("reference_type"), Payment.amount.label("amount"),
            Payment.payment_method.label("payment_method"), User.name.label("received_by"),
        ).outerjoin(Sale, and_(Sale.business_id == business_id, cast(Sale.id, String) == Payment.reference_id, Payment.reference_type.in_((PaymentReferenceType.SALE, PaymentReferenceType.CUSTOMER_PAYMENT, PaymentReferenceType.REFUND)))).outerjoin(
            Purchase, and_(Purchase.business_id == business_id, cast(Purchase.id, String) == Payment.reference_id, Payment.reference_type == PaymentReferenceType.PURCHASE)
        ).outerjoin(Customer, and_(Customer.id == Payment.customer_id, Customer.business_id == business_id)).outerjoin(
            Supplier, and_(Supplier.id == Purchase.supplier_id, Supplier.business_id == business_id)
        ).outerjoin(User, and_(User.id == Payment.received_by, User.business_id == business_id)).where(
            Payment.business_id == business_id, Payment.payment_date >= start_at, Payment.payment_date < end_at,
        )
        if filters.get("payment_method"):
            statement = statement.where(Payment.payment_method == filters["payment_method"])
        if filters.get("payment_type"):
            statement = statement.where(Payment.payment_type == filters["payment_type"])
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        payment_summary = select(
            Payment.payment_method, Payment.payment_type, Payment.reference_type, func.sum(Payment.amount),
        ).where(
            Payment.business_id == business_id,
            Payment.payment_date >= start_at,
            Payment.payment_date < end_at,
        )
        if filters.get("payment_method"):
            payment_summary = payment_summary.where(Payment.payment_method == filters["payment_method"])
        if filters.get("payment_type"):
            payment_summary = payment_summary.where(Payment.payment_type == filters["payment_type"])
        summary_rows = db.execute(payment_summary.group_by(
            Payment.payment_method, Payment.payment_type, Payment.reference_type,
        )).all()
        by_method = defaultdict(lambda: ZERO)
        by_type = defaultdict(lambda: ZERO)
        for method, payment_type, reference_type, amount in summary_rows:
            by_method[method.value] += amount or ZERO
            by_type[f"{payment_type.value}:{reference_type.value}"] += amount or ZERO
        rows = db.execute(statement.order_by(Payment.payment_date.desc()).offset((page - 1) * limit).limit(limit)).mappings().all()
        return {"items": [dict(row) for row in rows], "total": total, "summary": {
            "by_method": dict(by_method), "by_type": dict(by_type),
            "total": money(sum(by_method.values(), ZERO)),
        }}

    raise ValueError("Unknown report")


def financial_summary(db: Session, business_id, start: date, end: date) -> dict:
    start_at, end_at = business_datetime_bounds(db, business_id, start, end)
    sales = db.execute(
        select(
            func.coalesce(func.sum(Sale.subtotal + Sale.tax), 0),
            func.coalesce(func.sum(Sale.discount), 0),
        ).where(
            Sale.business_id == business_id,
            Sale.sale_date >= start_at,
            Sale.sale_date < end_at,
            Sale.status != SaleStatus.VOIDED,
        )
    ).one()
    refunds = db.scalar(select(func.coalesce(func.sum(SaleReturn.refund_amount), 0)).join(
        Sale, and_(Sale.id == SaleReturn.sale_id, Sale.business_id == business_id)
    ).where(
        SaleReturn.business_id == business_id,
        SaleReturn.created_at >= start_at,
        SaleReturn.created_at < end_at,
    )) or ZERO
    sold_cogs = db.scalar(
        select(func.coalesce(func.sum(SaleItem.cogs), 0))
        .join(Sale, and_(Sale.id == SaleItem.sale_id, Sale.business_id == business_id))
        .where(
            SaleItem.business_id == business_id,
            Sale.sale_date >= start_at,
            Sale.sale_date < end_at,
            Sale.status != SaleStatus.VOIDED,
        )
    ) or ZERO
    returned_cogs = db.scalar(
        select(func.coalesce(func.sum(SaleReturnItem.cogs), 0))
        .join(SaleReturn, and_(SaleReturn.id == SaleReturnItem.return_id, SaleReturn.business_id == business_id))
        .where(
            SaleReturnItem.business_id == business_id,
            SaleReturn.created_at >= start_at,
            SaleReturn.created_at < end_at,
        )
    ) or ZERO
    expense_total = db.scalar(select(func.coalesce(func.sum(Expense.amount), 0)).where(
        Expense.business_id == business_id, _date_filter(Expense.expense_date, start, end),
    )) or ZERO
    revenue, discounts = sales
    cogs = money(sold_cogs - returned_cogs)
    net_revenue = money(revenue - discounts - refunds)
    gross_profit = money(net_revenue - cogs)
    net_profit = money(gross_profit - expense_total)
    expense_rows = db.execute(select(ExpenseCategory.name, func.sum(Expense.amount)).join(
        ExpenseCategory, and_(ExpenseCategory.id == Expense.category_id, ExpenseCategory.business_id == business_id)
    ).where(Expense.business_id == business_id, _date_filter(Expense.expense_date, start, end)).group_by(ExpenseCategory.name).order_by(func.sum(Expense.amount).desc())).all()
    payment_rows = db.execute(select(Payment.payment_method, func.sum(Payment.amount)).where(
        Payment.business_id == business_id, Payment.payment_type == PaymentType.INCOMING,
        Payment.payment_date >= start_at, Payment.payment_date < end_at,
    ).group_by(Payment.payment_method)).all()
    previous_days = (end - start).days + 1
    previous_start = start - timedelta(days=previous_days)
    previous_end = start - timedelta(days=1)
    comparison = None
    if previous_days <= MAX_RANGE_DAYS:
        prev = _financial_totals_only(db, business_id, previous_start, previous_end)
        comparison = prev
    return {
        "period": {"start_date": start.isoformat(), "end_date": end.isoformat()},
        "revenue": revenue or ZERO,
        "discounts": discounts or ZERO,
        "returns": refunds,
        "net_revenue": net_revenue,
        "cogs": cogs,
        "gross_profit": gross_profit,
        "operating_expenses": expense_total,
        "net_profit": net_profit,
        "gross_margin_percent": _percent(gross_profit, net_revenue),
        "net_margin_percent": _percent(net_profit, net_revenue),
        "expense_breakdown": [{"category": name, "amount": amount} for name, amount in expense_rows],
        "payment_methods": [{"method": method.value, "amount": amount or ZERO} for method, amount in payment_rows],
        "comparison": comparison,
    }


def _financial_totals_only(db, business_id, start, end):
    start_at, end_at = business_datetime_bounds(db, business_id, start, end)
    row = db.execute(select(
        func.coalesce(func.sum(Sale.subtotal + Sale.tax), 0),
        func.coalesce(func.sum(Sale.discount), 0),
    ).where(Sale.business_id == business_id, Sale.sale_date >= start_at, Sale.sale_date < end_at, Sale.status != SaleStatus.VOIDED)).one()
    refunds = db.scalar(select(func.coalesce(func.sum(SaleReturn.refund_amount), 0)).join(
        Sale, and_(Sale.id == SaleReturn.sale_id, Sale.business_id == business_id)
    ).where(SaleReturn.business_id == business_id, SaleReturn.created_at >= start_at, SaleReturn.created_at < end_at)) or ZERO
    sold_cogs = db.scalar(
        select(func.coalesce(func.sum(SaleItem.cogs), 0))
        .join(Sale, and_(Sale.id == SaleItem.sale_id, Sale.business_id == business_id))
        .where(SaleItem.business_id == business_id, Sale.sale_date >= start_at, Sale.sale_date < end_at, Sale.status != SaleStatus.VOIDED)
    ) or ZERO
    returned_cogs = db.scalar(
        select(func.coalesce(func.sum(SaleReturnItem.cogs), 0))
        .join(SaleReturn, and_(SaleReturn.id == SaleReturnItem.return_id, SaleReturn.business_id == business_id))
        .where(SaleReturnItem.business_id == business_id, SaleReturn.created_at >= start_at, SaleReturn.created_at < end_at)
    ) or ZERO
    expenses = db.scalar(select(func.coalesce(func.sum(Expense.amount), 0)).where(
        Expense.business_id == business_id, _date_filter(Expense.expense_date, start, end),
    )) or ZERO
    net_revenue = money(row[0] - row[1] - refunds)
    cogs = money(sold_cogs - returned_cogs)
    gross_profit = money(net_revenue - cogs)
    net_profit = money(gross_profit - expenses)
    return {"revenue": row[0], "gross_profit": gross_profit, "net_profit": net_profit, "operating_expenses": expenses}


def daily_closing(db: Session, business_id, day: date) -> dict:
    start_at, end_at = business_datetime_bounds(db, business_id, day, day)
    sales = db.execute(select(
        func.coalesce(func.sum(Sale.subtotal - Sale.discount + Sale.tax), 0),
        func.count(Sale.id),
    ).where(Sale.business_id == business_id, Sale.sale_date >= start_at, Sale.sale_date < end_at, Sale.status != SaleStatus.VOIDED)).one()
    returns = db.scalar(select(func.coalesce(func.sum(SaleReturn.refund_amount), 0)).where(
        SaleReturn.business_id == business_id, SaleReturn.created_at >= start_at, SaleReturn.created_at < end_at,
    )) or ZERO
    sold_cogs = db.scalar(select(func.coalesce(func.sum(SaleItem.cogs), 0)).join(
        Sale, and_(Sale.id == SaleItem.sale_id, Sale.business_id == business_id)
    ).where(SaleItem.business_id == business_id, Sale.sale_date >= start_at, Sale.sale_date < end_at, Sale.status != SaleStatus.VOIDED)) or ZERO
    returned_cogs = db.scalar(select(func.coalesce(func.sum(SaleReturnItem.cogs), 0)).join(
        SaleReturn, and_(SaleReturn.id == SaleReturnItem.return_id, SaleReturn.business_id == business_id)
    ).where(SaleReturnItem.business_id == business_id, SaleReturn.created_at >= start_at, SaleReturn.created_at < end_at)) or ZERO
    cogs = money(sold_cogs - returned_cogs)
    expenses = db.scalar(select(func.coalesce(func.sum(Expense.amount), 0)).where(
        Expense.business_id == business_id, Expense.expense_date == day,
    )) or ZERO
    purchases = db.scalar(select(func.coalesce(func.sum(Purchase.total_amount), 0)).where(
        Purchase.business_id == business_id, Purchase.purchase_date == day,
    )) or ZERO
    payment_rows = db.execute(select(Payment.payment_method, Payment.reference_type, Payment.payment_type, func.sum(Payment.amount)).where(
        Payment.business_id == business_id,
        Payment.payment_date >= start_at, Payment.payment_date < end_at,
    ).group_by(Payment.payment_method, Payment.reference_type, Payment.payment_type)).all()
    method_totals = defaultdict(lambda: ZERO)
    customer_cash = ZERO
    refunds = ZERO
    supplier_cash = ZERO
    for method, reference, payment_type, amount in payment_rows:
        amount = amount or ZERO
        if payment_type == PaymentType.OUTGOING:
            if method.value == "cash" and reference == PaymentReferenceType.REFUND:
                refunds += amount
            elif method.value == "cash" and reference == PaymentReferenceType.PURCHASE:
                supplier_cash += amount
            continue
        if reference == PaymentReferenceType.CUSTOMER_PAYMENT:
            if method.value == "cash":
                customer_cash += amount
            continue
        method_totals[method.value] += amount
    existing = db.scalar(select(DailyClosing).where(
        DailyClosing.business_id == business_id, DailyClosing.closing_date == day,
    ))
    cash_expenses = db.scalar(select(func.coalesce(func.sum(Expense.amount), 0)).where(
        Expense.business_id == business_id, Expense.expense_date == day,
        Expense.payment_method == "cash",
    )) or ZERO
    opening_cash = existing.opening_cash if existing else ZERO
    expected_cash = money(opening_cash + method_totals["cash"] + customer_cash - refunds - supplier_cash - cash_expenses)
    closer_name = db.scalar(select(User.name).where(
        User.id == existing.closed_by, User.business_id == business_id,
    )) if existing and existing.closed_by else None
    net_sales = money(sales[0] - returns)
    gross_profit = money(net_sales - cogs)
    return {
        "date": day.isoformat(), "transactions": sales[1], "total_sales": sales[0] or ZERO,
        "total_returns": returns, "net_sales": net_sales,
        "opening_cash": opening_cash,
        "cash_sales": method_totals["cash"],
        "customer_payments": customer_cash,
        "supplier_payments": supplier_cash,
        "refunds": refunds,
        "expected_closing_cash": expected_cash,
        "actual_closing_cash": existing.actual_closing_cash if existing else None,
        "variance": money(existing.actual_closing_cash - expected_cash) if existing else None,
        "closed_by": closer_name,
        "closed_at": existing.closed_at.isoformat() if existing else None,
        "notes": existing.notes if existing else None,
        "cash_received": method_totals["cash"],
        "card_received": method_totals["card"],
        "bank_transfer": method_totals["bank_transfer"],
        "mobile_wallet": method_totals["mobile_wallet"],
        "other_payments": method_totals["other"],
        "purchases": purchases, "expenses": expenses, "cogs": cogs,
        "gross_profit": gross_profit, "net_profit": money(gross_profit - expenses),
    }


def dashboard_analytics(db: Session, business_id, start: date, end: date) -> dict:
    start_at, end_at = business_datetime_bounds(db, business_id, start, end)
    sale_day = business_local_date_expression(db, business_id, Sale.sale_date)
    return_day_expr = business_local_date_expression(db, business_id, SaleReturn.created_at)
    days = (end - start).days + 1
    if days > MAX_RANGE_DAYS:
        raise ValueError(f"Date range cannot exceed {MAX_RANGE_DAYS} days")
    sales = db.execute(
        select(
            sale_day,
            func.sum(Sale.subtotal - Sale.discount + Sale.tax),
        )
        .where(
            Sale.business_id == business_id,
            Sale.sale_date >= start_at,
            Sale.sale_date < end_at,
            Sale.status != SaleStatus.VOIDED,
        )
        .group_by(sale_day)
    ).all()
    sales_by_day = defaultdict(lambda: {"sales": ZERO, "returns": ZERO, "cogs": ZERO, "returned_cogs": ZERO})
    for sold_day, amount in sales:
        sales_by_day[str(sold_day)[:10]]["sales"] = amount or ZERO
    sale_cogs = db.execute(
        select(sale_day, func.sum(SaleItem.cogs))
        .join(Sale, and_(Sale.id == SaleItem.sale_id, Sale.business_id == business_id))
        .where(
            SaleItem.business_id == business_id,
            Sale.sale_date >= start_at,
            Sale.sale_date < end_at,
            Sale.status != SaleStatus.VOIDED,
        )
        .group_by(sale_day)
    ).all()
    for sold_day, amount in sale_cogs:
        sales_by_day[str(sold_day)[:10]]["cogs"] = amount or ZERO
    return_rows = db.execute(
        select(return_day_expr, func.sum(SaleReturn.refund_amount))
        .where(
            SaleReturn.business_id == business_id,
            SaleReturn.created_at >= start_at,
            SaleReturn.created_at < end_at,
        )
        .group_by(return_day_expr)
    ).all()
    for return_day, amount in return_rows:
        sales_by_day[str(return_day)[:10]]["returns"] = amount or ZERO
    return_cost_rows = db.execute(
        select(return_day_expr, func.sum(SaleReturnItem.cogs))
        .join(SaleReturn, and_(SaleReturn.id == SaleReturnItem.return_id, SaleReturn.business_id == business_id))
        .where(
            SaleReturnItem.business_id == business_id,
            SaleReturn.created_at >= start_at,
            SaleReturn.created_at < end_at,
        )
        .group_by(return_day_expr)
    ).all()
    for return_day, amount in return_cost_rows:
        sales_by_day[str(return_day)[:10]]["returned_cogs"] = amount or ZERO
    expenses = db.execute(select(Expense.expense_date, func.sum(Expense.amount)).where(
        Expense.business_id == business_id, _date_filter(Expense.expense_date, start, end),
    ).group_by(Expense.expense_date)).all()
    expense_by_day = {day.isoformat(): amount for day, amount in expenses}
    purchases = db.execute(select(Purchase.purchase_date, func.sum(Purchase.total_amount)).where(
        Purchase.business_id == business_id, _date_filter(Purchase.purchase_date, start, end),
    ).group_by(Purchase.purchase_date)).all()
    purchase_by_day = {day.isoformat(): amount for day, amount in purchases}
    sales_trend = []
    expense_trend = []
    purchase_trend = []
    for offset in range(days):
        day = (start + timedelta(days=offset)).isoformat()
        record = sales_by_day[day]
        net_sales = money(record["sales"] - record["returns"])
        net_cogs = money(record["cogs"] - record["returned_cogs"])
        gross_profit = money(net_sales - net_cogs)
        net_profit = money(gross_profit - expense_by_day.get(day, ZERO))
        sales_trend.append({
            "date": day,
            "sales": net_sales,
            "cogs": net_cogs,
            "gross_profit": gross_profit,
            "expenses": money(expense_by_day.get(day, ZERO)),
            "net_profit": net_profit,
        })
        expense_trend.append({"date": day, "expenses": money(expense_by_day.get(day, ZERO))})
        purchase_trend.append({"date": day, "purchases": money(purchase_by_day.get(day, ZERO))})
    product_profit = _profitability_rows(db, business_id, start, end)
    category_profit = _profitability_rows(db, business_id, start, end, by_category=True)
    payment_rows = db.execute(select(Payment.payment_method, func.sum(Payment.amount)).where(
        Payment.business_id == business_id, Payment.payment_type == PaymentType.INCOMING,
        Payment.payment_date >= start_at, Payment.payment_date < end_at,
    ).group_by(Payment.payment_method)).all()
    top_customers = db.execute(select(Customer.name, func.sum(Sale.total_amount)).join(
        Sale, and_(Sale.customer_id == Customer.id, Sale.business_id == business_id)
    ).where(Customer.business_id == business_id, Sale.sale_date >= start_at, Sale.sale_date < end_at, Sale.status != SaleStatus.VOIDED)
        .group_by(Customer.id, Customer.name).order_by(func.sum(Sale.total_amount).desc()).limit(5)).all()
    low_stock = db.execute(select(Product.name, Product.sku, Product.current_stock, Product.minimum_stock).where(
        Product.business_id == business_id, Product.is_active.is_(True),
        Product.current_stock <= Product.minimum_stock,
    ).order_by(Product.current_stock, Product.name).limit(10)).all()
    financials = financial_summary(db, business_id, start, end)
    inventory = db.execute(
        select(
            func.coalesce(func.sum(Product.current_stock * Product.cost_price), 0),
            func.coalesce(func.sum(Product.current_stock * Product.selling_price), 0),
            func.count(Product.id),
            func.sum(case((Product.current_stock <= Product.minimum_stock, 1), else_=0)),
            func.sum(case((Product.current_stock <= 0, 1), else_=0)),
        ).where(Product.business_id == business_id, Product.is_active.is_(True))
    ).one()
    customer_outstanding = db.scalar(
        select(func.coalesce(func.sum(Customer.current_balance), 0)).where(Customer.business_id == business_id)
    ) or ZERO
    supplier_outstanding = db.scalar(
        select(func.coalesce(func.sum(Purchase.outstanding_amount), 0)).where(Purchase.business_id == business_id)
    ) or ZERO
    return {
        "sales_trend": sales_trend,
        "profit_trend": [{"date": row["date"], "profit": row["net_profit"]} for row in sales_trend],
        "expense_trend": expense_trend,
        "revenue_vs_expenses": [{"date": row["date"], "revenue": row["sales"], "expenses": expense_trend[index]["expenses"]} for index, row in enumerate(sales_trend)],
        "top_products": product_profit[:5],
        "top_categories": category_profit[:5],
        "payment_methods": [{"method": method.value, "amount": amount or ZERO} for method, amount in payment_rows],
        "low_stock_products": [{"name": name, "sku": sku, "current_stock": stock, "minimum_stock": minimum} for name, sku, stock, minimum in low_stock],
        "customer_sales": [{"customer": name, "sales": amount} for name, amount in top_customers],
        "purchase_trend": [{"date": (start + timedelta(days=i)).isoformat(), "purchases": money(purchase_by_day.get((start + timedelta(days=i)).isoformat(), ZERO))} for i in range(days)],
        "financials": financials,
        "inventory": {
            "cost_value": inventory[0] or ZERO,
            "retail_value": inventory[1] or ZERO,
            "potential_gross_profit": money((inventory[1] or ZERO) - (inventory[0] or ZERO)),
            "products": inventory[2] or 0,
            "low_stock": inventory[3] or 0,
            "out_of_stock": inventory[4] or 0,
        },
        "customer_outstanding": customer_outstanding,
        "supplier_outstanding": supplier_outstanding,
    }


def make_excel_report(business_name: str, title: str, start: date, end: date, fields: list[str], rows: list[dict], summary: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(fields)))
    sheet.cell(1, 1, business_name).font = Font(size=16, bold=True, color="25324B")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, len(fields)))
    sheet.cell(2, 1, title).font = Font(size=13, bold=True)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(1, len(fields)))
    sheet.cell(3, 1, f"Period: {start.isoformat()} to {end.isoformat()} | Generated: {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC}")
    headers = [field.replace("_", " ").title() for field in fields]
    header_row = 5
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="35445D")
        cell.alignment = Alignment(horizontal="center")
    for row_index, row in enumerate(rows, start=header_row + 1):
        for column, field in enumerate(fields, start=1):
            value = row.get(field)
            if hasattr(value, "value"):
                value = value.value
            if isinstance(value, (date, datetime)):
                value = value.isoformat()
            cell = sheet.cell(row_index, column, value)
            if isinstance(value, (Decimal, int, float)):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
    summary_row = header_row + len(rows) + 2
    for key, value in summary.items():
        if isinstance(value, (int, float, Decimal)):
            sheet.cell(summary_row, 1, key.replace("_", " ").title()).font = Font(bold=True)
            cell = sheet.cell(summary_row, 2, float(value) if isinstance(value, Decimal) else value)
            cell.number_format = '#,##0.00;[Red]-#,##0.00'
            summary_row += 1
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:{get_column_letter(max(1, len(fields)))}{header_row + len(rows)}"
    for index, field in enumerate(fields, start=1):
        values = [field] + [str(row.get(field, "")) for row in rows[:200]]
        width = min(max(max(len(value) for value in values) + 2, 12), 34)
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pdf_report(business_name: str, title: str, start: date, end: date, fields: list[str], rows: list[dict], summary: dict) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(letter), rightMargin=0.4 * inch, leftMargin=0.4 * inch, topMargin=0.45 * inch, bottomMargin=0.45 * inch)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(business_name, styles["Title"]),
        Paragraph(title, styles["Heading2"]),
        Paragraph(f"Period: {start.isoformat()} to {end.isoformat()} · Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC}", styles["Normal"]),
        Spacer(1, 14),
    ]
    table_rows = [[field.replace("_", " ").title() for field in fields]]
    for row in rows:
        table_rows.append([str(row.get(field, "") if row.get(field) is not None else "—") for field in fields])
    table = Table(table_rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#35445d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#dfe3e9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f8fa")]),
    ]))
    story.append(table)
    if summary:
        story.extend([Spacer(1, 12), Paragraph("Summary", styles["Heading3"])])
        summary_rows = [[key.replace("_", " ").title(), f"{value:,.2f}" if isinstance(value, (int, float, Decimal)) else str(value)] for key, value in summary.items() if isinstance(value, (int, float, Decimal, str))]
        summary_table = Table(summary_rows, colWidths=[2.5 * inch, 1.5 * inch])
        summary_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#dfe3e9")), ("ALIGN", (1, 0), (1, -1), "RIGHT")]))
        story.append(summary_table)
    document.build(story)
    return output.getvalue()
