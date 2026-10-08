from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook


def register(client, business_name, email):
    response = client.post("/auth/register", json={
        "business_name": business_name,
        "owner_name": "Day Three Owner",
        "email": email,
        "password": "Strong-test-password-123!",
        "confirm_password": "Strong-test-password-123!",
    })
    assert response.status_code == 201, response.text
    token = response.json()["access_token"]
    return {"Authorization": "Bearer " + token}


def test_expenses_reports_exports_and_tenant_isolation(client):
    owner_a = register(client, "Day Three Store A", "owner-a@day3.test")
    owner_b = register(client, "Day Three Store B", "owner-b@day3.test")
    today = datetime.now(timezone.utc).date().isoformat()

    categories_response = client.get("/expenses/categories", headers=owner_a)
    assert categories_response.status_code == 200, categories_response.text
    category = next(item for item in categories_response.json() if item["name"] == "Rent")

    supplier_response = client.post("/suppliers", headers=owner_a, json={"name": "Day Three Supplier"})
    assert supplier_response.status_code == 201, supplier_response.text
    product_response = client.post("/products", headers=owner_a, json={
        "sku": "DAY3-WIDGET",
        "name": "Day Three Widget",
        "unit": "pcs",
        "cost_price": "60.00",
        "selling_price": "100.00",
        "minimum_stock": "2",
        "opening_stock": "10",
    })
    assert product_response.status_code == 201, product_response.text
    product_id = product_response.json()["id"]

    customer_response = client.post("/customers", headers=owner_a, json={"name": "Day Three Customer"})
    assert customer_response.status_code == 201, customer_response.text
    sale_response = client.post("/sales", headers=owner_a, json={
        "customer_id": customer_response.json()["id"],
        "paid_amount": "500.00",
        "items": [{"product_id": product_id, "quantity": "5"}],
    })
    assert sale_response.status_code == 201, sale_response.text
    sale = sale_response.json()
    assert Decimal(str(sale["cogs_amount"])) == Decimal("300.00")
    assert Decimal(str(sale["gross_profit"])) == Decimal("200.00")

    expense_response = client.post("/expenses", headers=owner_a, json={
        "category_id": category["id"],
        "description": "Monthly rent",
        "amount": "50.00",
        "expense_date": today,
        "payment_method": "bank_transfer",
        "reference_number": "RENT-1",
        "notes": "Test expense",
    })
    assert expense_response.status_code == 201, expense_response.text
    expense_id = expense_response.json()["id"]
    assert Decimal(str(client.get("/expenses", headers=owner_a).json()["total_amount"])) == Decimal("50.00")

    pnl = client.get("/reports/profit-and-loss", headers=owner_a, params={"start_date": today, "end_date": today})
    assert pnl.status_code == 200, pnl.text
    assert Decimal(str(pnl.json()["gross_profit"])) == Decimal("200.00")
    assert Decimal(str(pnl.json()["net_profit"])) == Decimal("150.00")

    profitability = client.get("/reports/product-profitability", headers=owner_a, params={"start_date": today, "end_date": today})
    assert profitability.status_code == 200, profitability.text
    product_row = next(row for row in profitability.json()["items"] if row["product"] == "Day Three Widget")
    assert Decimal(str(product_row["revenue"])) == Decimal("500.00")
    assert Decimal(str(product_row["cogs"])) == Decimal("300.00")
    assert Decimal(str(product_row["gross_profit"])) == Decimal("200.00")

    returned = client.post(f"/sales/{sale['id']}/returns", headers=owner_a, json={
        "items": [{"sale_item_id": sale["items"][0]["id"], "quantity": "2"}],
        "payment_method": "cash",
    })
    assert returned.status_code == 201, returned.text
    adjusted_pnl = client.get("/reports/profit-and-loss", headers=owner_a, params={"start_date": today, "end_date": today})
    assert adjusted_pnl.status_code == 200, adjusted_pnl.text
    assert Decimal(str(adjusted_pnl.json()["net_revenue"])) == Decimal("300.00")
    assert Decimal(str(adjusted_pnl.json()["cogs"])) == Decimal("180.00")
    assert Decimal(str(adjusted_pnl.json()["gross_profit"])) == Decimal("120.00")
    assert Decimal(str(adjusted_pnl.json()["net_profit"])) == Decimal("70.00")

    inventory = client.get("/reports/inventory", headers=owner_a)
    assert inventory.status_code == 200, inventory.text
    assert Decimal(str(inventory.json()["summary"]["inventory_cost_value"])) == Decimal("420.00")

    closing = client.get("/reports/daily-closing", headers=owner_a, params={"report_date": today})
    assert closing.status_code == 200, closing.text
    assert Decimal(str(closing.json()["summary"]["net_profit"])) == Decimal("70.00")
    payment_report = client.get("/reports/payments", headers=owner_a, params={"start_date": today, "end_date": today})
    assert payment_report.status_code == 200, payment_report.text
    assert Decimal(str(payment_report.json()["summary"]["by_type"]["incoming:sale"])) == Decimal("500.00")
    assert Decimal(str(payment_report.json()["summary"]["by_type"]["outgoing:refund"])) == Decimal("200.00")
    assert {item["reference_type"] for item in payment_report.json()["items"]} >= {"sale", "refund"}

    close_response = client.post("/reports/daily-closing", headers=owner_a, json={
        "report_date": today,
        "opening_cash": "0.00",
        "actual_closing_cash": "300.00",
        "notes": "Verified test close",
    })
    assert close_response.status_code == 201, close_response.text
    assert Decimal(str(close_response.json()["summary"]["expected_closing_cash"])) == Decimal("300.00")
    assert Decimal(str(close_response.json()["summary"]["variance"])) == Decimal("0.00")
    duplicate_close = client.post("/reports/daily-closing", headers=owner_a, json={
        "report_date": today,
        "actual_closing_cash": "300.00",
    })
    assert duplicate_close.status_code == 409

    for report_name in ("sales", "inventory", "product-profitability", "profit-and-loss", "daily-closing"):
        exported = client.get(
            f"/reports/{report_name}/export.xlsx",
            headers=owner_a,
            params={"start_date": today, "end_date": today},
        )
        assert exported.status_code == 200, f"{report_name}: {exported.text}"
        assert exported.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        assert workbook.active.max_row > 4
        workbook.close()

    for report_name in ("sales", "inventory", "profit-and-loss", "daily-closing"):
        exported = client.get(
            f"/reports/{report_name}/export.pdf",
            headers=owner_a,
            params={"start_date": today, "end_date": today},
        )
        assert exported.status_code == 200, f"{report_name}: {exported.text}"
        assert exported.headers["content-type"].startswith("application/pdf")
        assert exported.content.startswith(b"%PDF-")

    assert client.get("/expenses", headers=owner_b).json()["total"] == 0
    assert client.get("/reports/sales", headers=owner_b, params={"start_date": today, "end_date": today}).json()["total"] == 0
    assert client.get("/reports/sales/export.xlsx", headers=owner_b, params={"start_date": today, "end_date": today}).status_code == 200
    assert client.get("/audit-logs", headers=owner_b).json()["total"] == 0
    assert client.get("/dashboard/overview", headers=owner_b).json()["today"]["transactions"] == 0
    assert client.get("/business", headers=owner_b).json()["business_name"] == "Day Three Store B"
    assert client.get(f"/expenses/{expense_id}", headers=owner_b).status_code == 404

    updated = client.put(f"/expenses/{expense_id}", headers=owner_a, json={"amount": "55.00"})
    assert updated.status_code == 200, updated.text
    deleted = client.delete(f"/expenses/{expense_id}", headers=owner_a)
    assert deleted.status_code == 204
    assert client.get(f"/expenses/{expense_id}", headers=owner_a).status_code == 404

    audits = client.get("/audit-logs", headers=owner_a)
    assert audits.status_code == 200, audits.text
    assert any(row["entity_type"] == "expense" for row in audits.json()["items"])
